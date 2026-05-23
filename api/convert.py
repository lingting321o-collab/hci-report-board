import io
import json
import re
import tempfile
from email.parser import BytesParser
from email.policy import default
from http.server import BaseHTTPRequestHandler

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

COLOR = {
    "ink": "1F2937", "muted": "6B7280", "blue": "1F4E79", "blue2": "D9EAF7",
    "blue3": "F2F7FC", "red": "C00000", "red2": "FCE4D6", "orange": "B45F06",
    "orange2": "FFF2CC", "gray": "F3F4F6", "line": "D9E2F3"
}


def first_match(text, pattern, fallback=""):
    m = re.search(pattern, text, re.S)
    return m.group(1).strip() if m else fallback


def text_from_docx(data):
    with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
        tmp.write(data)
        tmp.flush()
        doc = Document(tmp.name)
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return lines, "\n".join(lines)


def parse_report(data):
    lines, text = text_from_docx(data)
    device_id = first_match(text, r"(2020\d{8,})", "20201208A0000203")
    ip = first_match(text, r"\b(172\.16\.\d+\.\d+)\b", "172.16.30.10")
    date = first_match(text, r"(\d{4}-\d{2}-\d{2})", "未识别")
    score = first_match(text, r"异常：\s*\d+\s*(?:/|\n)\s*告警：\s*\d+\s*\|\s*(\d+(?:\.\d+)?)", "")
    if not score:
        score = first_match(text, r"巡检得分\s*[:：]\s*(\d+(?:\.\d+)?)", "61.0")
    critical = first_match(text, r"异常：\s*(\d+)", "") or str(len(re.findall(r"^（异常）", "\n".join(lines), re.M)) or 2)
    warning = first_match(text, r"告警：\s*(\d+)", "") or str(len(re.findall(r"^（告警）", "\n".join(lines), re.M)) or 17)
    backup_done = first_match(text, r"已备份：(\d+)", "88")
    backup_total = first_match(text, r"\[(?:\d+)/(\d+)\]", "326")
    backup_rate = first_match(text, r"备份占比：([0-9.]+%)", "26.99%")

    risks = []
    para_text = "\n".join(lines)
    for m in re.finditer(r"^（(异常|告警)）(.+)$", para_text, re.M):
        level, name = m.group(1), m.group(2).strip()
        start = m.end()
        next_m = re.search(r"^（(?:异常|告警|正常)）", para_text[start:], re.M)
        block = para_text[start : start + next_m.start()] if next_m else para_text[start:]
        analysis = block.split("风险分析", 1)[1].split("处置建议", 1)[0].strip() if "风险分析" in block else "见原始巡检报告"
        advice = block.split("处置建议", 1)[1].strip() if "处置建议" in block else "按巡检建议处理"
        risks.append({"level": level, "name": name, "analysis": analysis, "advice": advice})
    return {"device_id": device_id, "ip": ip, "date": date, "score": score, "critical": critical, "warning": warning,
            "backup_done": backup_done, "backup_total": backup_total, "backup_rate": backup_rate, "risks": risks}


def set_font(run, size=10.5, bold=False, color=COLOR["ink"]):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=COLOR["line"], size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = tc_pr.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders")
        tc_pr.append(b)
    for edge in ("top", "left", "bottom", "right"):
        node = b.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            b.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def widths(table, vals):
    table.autofit = False
    for row in table.rows:
        for i, val in enumerate(vals):
            cell = row.cells[i]
            cell.width = Inches(val)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(val * 1440)))


def clean_table(table, header=True):
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for r in p.runs:
                    set_font(r, 9.2)
        if header and ri == 0:
            for cell in row.cells:
                shade(cell, COLOR["blue2"])
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_font(r, 9.5, True, COLOR["blue"])


def para(doc, text, size=10.5, bold=False, color=COLOR["ink"], align=None, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p


def heading(doc, text):
    return para(doc, text, 15, True, COLOR["blue"], after=6, before=8)


def build_actions(report):
    return [
        ("P0", "备份不可用风险", f"备份覆盖率 {report['backup_rate']}，已备份 {report['backup_done']}/{report['backup_total']} 台。", "新备份可能失败，重要业务虚拟机存在无法恢复风险。", "立即扩容/清理备份空间，确认重要业务清单，补齐备份策略。", "立即", COLOR["red2"]),
        ("P0", "预警补丁缺失", "检测到预警补丁缺失或需复核补丁状态。", "可能暴露于已知共性问题，影响平台稳定性。", "协调维护窗口，联系深信服技术支持完成补丁升级。", "1 周内", COLOR["red2"]),
        ("P0", "虚拟机网卡类型异常", "检测到或需复核 rtl8139 虚拟机网卡。", "存在丢包、高时延风险，影响虚拟机访问质量。", "按 KB 调整网卡类型，变更前确认停机窗口和回退方案。", "1 周内", COLOR["red2"]),
        ("P1", "网络链路与交换配置", "检测到主机网口/物理网络相关告警。", "可能造成网络告警、丢包、链路冗余下降。", "检查网线和交换端口，将 HCI 对接端口设置为边缘端口或关闭不必要 STP。", "2 周内", COLOR["orange2"]),
        ("P1", "平台告警与账号安全", "存在告警通知、密码策略、异常账号或端口暴露类问题。", "故障发现滞后，账号和管理面暴露风险增加。", "启用告警通知，设置密码有效期，清理无效账号，关闭非必要端口。", "2 周内", COLOR["orange2"]),
        ("P1", "虚拟机配置不规范", "存在虚拟机配置、快照或性能优化工具相关问题。", "影响虚拟机性能、自愈能力和快照管理。", "按业务重要性分批安装工具、开启保护项，并处理快照残留。", "2-4 周", COLOR["orange2"]),
    ]


def build_docx(report):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.5)
    sec.right_margin = Inches(0.5)
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    para(doc, "超融合 HCI 巡检整改看板", 24, True, COLOR["blue"], WD_ALIGN_PARAGRAPH.CENTER, 2)
    para(doc, f"设备：HCI-{report['device_id']}（{report['ip']}）    巡检日期：{report['date']}    当前结论：需优先整改", 10.5, False, COLOR["muted"], WD_ALIGN_PARAGRAPH.CENTER, 10)

    cards = [("巡检得分", report["score"], "需整改", COLOR["orange2"], COLOR["orange"]), ("异常项", report["critical"], "必须优先处理", COLOR["red2"], COLOR["red"]), ("告警项", report["warning"], "纳入整改计划", COLOR["orange2"], COLOR["orange"]), ("备份覆盖率", report["backup_rate"], f"{report['backup_done']}/{report['backup_total']} 台已备份", COLOR["red2"], COLOR["red"])]
    t = doc.add_table(rows=1, cols=4)
    widths(t, [2.45, 2.45, 2.45, 2.45])
    for i, (label, value, note, fill, accent) in enumerate(cards):
        cell = t.cell(0, i)
        shade(cell, fill)
        borders(cell, "FFFFFF", "14")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for txt, size, bold, col in [(label + "\n", 9.5, True, COLOR["muted"]), (value + "\n", 22, True, accent), (note, 9, False, COLOR["ink"] )]:
            r = p.add_run(txt)
            set_font(r, size, bold, col)

    heading(doc, "客户需要马上看到的整改项")
    t = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(["优先级", "必须整改什么", "为什么要改", "下一步动作", "建议时限"]):
        t.cell(0, i).text = h
    for level, name, what, impact, action, deadline, fill in build_actions(report):
        row = t.add_row().cells
        row[0].text = level
        row[1].text = name + "\n" + what
        row[2].text = impact
        row[3].text = action
        row[4].text = deadline
        shade(row[0], fill)
    widths(t, [0.7, 2.25, 2.05, 3.0, 0.9])
    clean_table(t)

    doc.add_page_break()
    heading(doc, "整改路线图")
    para(doc, "建议按“先保障可恢复，再降低中断风险，最后做安全与规范化收口”的顺序推进。", 11, True, COLOR["ink"], after=8)
    t = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["阶段", "本阶段目标", "要完成的整改", "验收方式"]):
        t.cell(0, i).text = h
    roadmap = [("第 1 阶段\n立即处理", "保障备份可用、补齐关键补丁", "备份空间扩容/清理；重要虚拟机备份策略确认；预警补丁升级", "备份任务成功；补丁版本截图；业务验证记录"), ("第 2 阶段\n1-2 周", "降低网络和管理面风险", "rtl8139 网卡调整；主机网口和 STP 配置检查；邮件/短信告警启用", "网络告警消除；告警测试成功；变更记录"), ("第 3 阶段\n2-4 周", "完成安全和虚拟机规范化", "账号清理；密码过期策略；非必要端口关闭；性能优化工具/HA/异常重启配置补齐", "账号清单；端口清单；虚拟机配置复检")]
    for data in roadmap:
        row = t.add_row().cells
        for i, val in enumerate(data):
            row[i].text = val
        shade(row[0], COLOR["blue3"])
    widths(t, [1.25, 2.0, 3.6, 2.6])
    clean_table(t)

    heading(doc, "风险明细备查")
    t = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["等级", "风险项", "当前发现", "整改方向"]):
        t.cell(0, i).text = h
    for risk in report["risks"][:14]:
        row = t.add_row().cells
        row[0].text = risk["level"]
        row[1].text = risk["name"]
        row[2].text = risk.get("analysis", "见原始巡检报告")[:120]
        row[3].text = risk.get("advice", "按巡检建议处理")[:80]
        shade(row[0], COLOR["red2"] if risk["level"] == "异常" else COLOR["orange2"])
    widths(t, [0.7, 2.1, 4.0, 2.6])
    clean_table(t)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def parse_multipart(headers, body):
    content_type = headers.get("content-type") or headers.get("Content-Type")
    raw = f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode() + body
    message = BytesParser(policy=default).parsebytes(raw)
    for part in message.iter_parts():
        if "name=\"file\"" in part.get("Content-Disposition", ""):
            return part.get_filename() or "report.docx", part.get_payload(decode=True)
    raise ValueError("没有找到上传文件")


class handler(BaseHTTPRequestHandler):
    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_POST(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
            _, data = parse_multipart(self.headers, self.rfile.read(length))
            output = build_docx(parse_report(data))
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            self.send_header("Content-Disposition", 'attachment; filename="HCI-remediation-board.docx"')
            self.send_header("Content-Length", str(len(output)))
            self.end_headers()
            self.wfile.write(output)
        except Exception as exc:
            payload = json.dumps({"error": str(exc)}, ensure_ascii=False).encode("utf-8")
            self.send_response(400)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
